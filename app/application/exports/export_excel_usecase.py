from datetime import date
from io import BytesIO
from typing import Optional
import uuid

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.domain.repositories.category_repository import CategoryRepository
from app.domain.repositories.expense_repository import ExpenseRepository


class ExportExpensesExcelUseCase:
    def __init__(self, expense_repo: ExpenseRepository, category_repo: CategoryRepository):
        self.expense_repo = expense_repo
        self.category_repo = category_repo

    def execute(
        self,
        user_id: uuid.UUID,
        start_date: Optional[date] = None,
        end_date: Optional[date] = None,
    ) -> BytesIO:
        wb = openpyxl.Workbook()

        # Styles
        header_fill = PatternFill(start_color="3F51B5", end_color="3F51B5", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="1A237E")
        bold_font = Font(name="Calibri", size=11, bold=True)
        thin_border = Border(
            left=Side(style="thin", color="E0E0E0"),
            right=Side(style="thin", color="E0E0E0"),
            top=Side(style="thin", color="E0E0E0"),
            bottom=Side(style="thin", color="E0E0E0"),
        )
        total_fill = PatternFill(start_color="EDE7F6", end_color="EDE7F6", fill_type="solid")

        # -------------------------------------------------------------
        # Sheet 1: Expenses List
        # -------------------------------------------------------------
        ws_expenses = wb.active
        ws_expenses.title = "Expenses"

        # Title block
        ws_expenses["A1"] = "Expense Tracker — Detailed Transactions"
        ws_expenses["A1"].font = title_font
        date_range_str = f"Date Range: {start_date or 'Beginning'} to {end_date or 'Present'}"
        ws_expenses["A2"] = date_range_str
        ws_expenses["A2"].font = Font(italic=True, color="555555")

        headers = ["Date", "Category", "Amount", "Description", "Created At", "Type"]
        for col_num, header in enumerate(headers, 1):
            cell = ws_expenses.cell(row=4, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center" if col_num != 4 else "left", vertical="center")

        # Fetch expenses
        expenses, _ = self.expense_repo.list(
            user_id=user_id,
            start_date=start_date,
            end_date=end_date,
            limit=10000,
        )

        # Build category map
        categories = {cat.id: cat.name_key.capitalize() for cat in self.category_repo.list_all(active_only=False)}

        row_num = 5
        total_amount = 0.0
        for exp in expenses:
            category_name = categories.get(exp.category_id, "Unknown")
            amt = float(exp.amount)
            if exp.entry_type != "credit":
                total_amount += amt

            ws_expenses.cell(row=row_num, column=1, value=str(exp.expense_date)).alignment = Alignment(horizontal="center")
            ws_expenses.cell(row=row_num, column=2, value=category_name)
            amt_cell = ws_expenses.cell(row=row_num, column=3, value=amt)
            amt_cell.number_format = "#,##0.00"
            ws_expenses.cell(row=row_num, column=4, value=exp.description or "")
            ws_expenses.cell(row=row_num, column=5, value=exp.created_at.strftime("%Y-%m-%d %H:%M")).alignment = Alignment(horizontal="center")
            ws_expenses.cell(row=row_num, column=6, value=exp.entry_type)

            for col in range(1, 7):
                ws_expenses.cell(row=row_num, column=col).border = thin_border
            row_num += 1

        # Total row
        ws_expenses.cell(row=row_num, column=2, value="Total").font = bold_font
        tot_cell = ws_expenses.cell(row=row_num, column=3, value=total_amount)
        tot_cell.font = bold_font
        tot_cell.number_format = "#,##0.00"
        for col in range(1, 7):
            c = ws_expenses.cell(row=row_num, column=col)
            c.fill = total_fill
            c.border = thin_border

        # -------------------------------------------------------------
        # Sheet 2: Category Breakdown
        # -------------------------------------------------------------
        ws_cats = wb.create_sheet(title="Category Breakdown")
        ws_cats["A1"] = "Category Spending Summary"
        ws_cats["A1"].font = title_font

        cat_headers = ["Category", "Total Spent", "Count", "Percentage (%)"]
        for col_num, header in enumerate(cat_headers, 1):
            cell = ws_cats.cell(row=3, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        cat_totals = self.expense_repo.get_category_totals(user_id=user_id, start_date=start_date, end_date=end_date)
        cat_row = 4
        for item in cat_totals:
            ws_cats.cell(row=cat_row, column=1, value=item["name_key"].capitalize())
            amt_cell = ws_cats.cell(row=cat_row, column=2, value=item["total_amount"])
            amt_cell.number_format = "#,##0.00"
            ws_cats.cell(row=cat_row, column=3, value=item["count"]).alignment = Alignment(horizontal="center")
            pct_cell = ws_cats.cell(row=cat_row, column=4, value=item["percentage"] / 100)
            pct_cell.number_format = "0.0%"
            pct_cell.alignment = Alignment(horizontal="center")

            for col in range(1, 5):
                ws_cats.cell(row=cat_row, column=col).border = thin_border
            cat_row += 1

        # -------------------------------------------------------------
        # Auto-adjust column widths for all sheets
        # -------------------------------------------------------------
        for sheet in (ws_expenses, ws_cats):
            for col in sheet.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        val_str = str(cell.value)
                        if len(val_str) > max_len:
                            max_len = len(val_str)
                sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output

from datetime import date
from io import BytesIO
import openpyxl


def test_excel_export(client):
    # Seed categories
    seed_res = client.post("/api/v1/categories/seed")
    food_cat = next(c for c in seed_res.json() if c["name_key"] == "food")

    # Register user
    reg = client.post(
        "/api/v1/auth/register",
        json={"email": "export_test@example.com", "password": "Password123!"},
    )
    token = reg.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # Add expense
    client.post(
        "/api/v1/expenses",
        headers=headers,
        json={
            "category_id": food_cat["id"],
            "amount": 1250.00,
            "expense_date": str(date.today()),
            "description": "Monthly Dinner Outing",
        },
    )

    # Request Excel export
    export_res = client.get("/api/v1/exports/expenses.xlsx", headers=headers)
    assert export_res.status_code == 200
    assert (
        export_res.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment; filename=" in export_res.headers["content-disposition"]

    # Validate Excel content with openpyxl
    file_bytes = BytesIO(export_res.content)
    wb = openpyxl.load_workbook(file_bytes)
    assert "Expenses" in wb.sheetnames
    assert "Category Breakdown" in wb.sheetnames

    ws_expenses = wb["Expenses"]
    # Check header
    assert ws_expenses.cell(row=4, column=1).value == "Date"
    assert ws_expenses.cell(row=4, column=2).value == "Category"
    # Check data row
    assert ws_expenses.cell(row=5, column=2).value == "Food"
    assert ws_expenses.cell(row=5, column=3).value == 1250.00
